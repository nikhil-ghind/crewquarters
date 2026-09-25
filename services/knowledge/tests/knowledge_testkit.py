"""Knowledge test fixtures and generated documents. Registered with
``pytest_plugins = ["knowledge_testkit"]``; not a conftest, because every ``conftest.py``
shares one module name and would shadow the root conftest."""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from pathlib import Path

import httpx
import pytest
from sqlalchemy import create_engine
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from crewquarters_knowledge.config import KnowledgeSettings
from crewquarters_knowledge.embeddings import HashingEmbedder
from crewquarters_knowledge.main import create_app
from crewquarters_knowledge.models import Document, DocumentChunk, KnowledgeBase
from crewquarters_shared.config import Settings
from crewquarters_shared.db.base import Base


def make_pdf(pages: list[str]) -> bytes:
    """A minimal PDF with one Helvetica text line per page (an empty string gives a page
    with no text layer, like a scan)."""
    objects: list[bytes] = []
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def make_docx(heading: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.fixture(scope="session")
def person3_knowledge_tables(database_url: str) -> None:
    """Create the tables until their reviewed Alembic migration merges (then a no-op)."""
    engine = create_engine(database_url)
    Base.metadata.create_all(
        engine,
        tables=[KnowledgeBase.__table__, Document.__table__, DocumentChunk.__table__],
    )
    engine.dispose()


class KnowledgeHarness:
    def __init__(self, settings: KnowledgeSettings) -> None:
        self.settings = settings
        self.app = create_app(settings, embedder=HashingEmbedder(), run_worker=False)
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.app),
            base_url="http://knowledge",
            headers={
                "authorization": f"Bearer {settings.internal_service_token.get_secret_value()}"
            },
        )

    async def drain(self) -> None:
        """Run queued ingestion jobs to completion."""
        while await self.app.state.worker.run_once():
            pass

    async def create_kb(self, owner_id: uuid.UUID, name: str = "Handbook") -> str:
        resp = await self.client.post(
            "/internal/v1/knowledge-bases", json={"ownerId": str(owner_id), "name": name}
        )
        assert resp.status_code == 201, resp.text
        return str(resp.json()["id"])

    async def upload(self, kb_id: str, name: str, content: bytes) -> httpx.Response:
        return await self.client.post(
            f"/internal/v1/knowledge-bases/{kb_id}/documents", files={"file": (name, content)}
        )

    async def query(self, kb_id: str, text: str, **extra: object) -> dict[str, object]:
        resp = await self.client.post(
            f"/internal/v1/knowledge-bases/{kb_id}/query", json={"query": text, **extra}
        )
        assert resp.status_code == 200, resp.text
        body: dict[str, object] = resp.json()
        return body


@pytest.fixture
async def knowledge(
    settings: Settings,
    tmp_path: Path,
    sessions: async_sessionmaker[AsyncSession],
    person3_knowledge_tables: None,
) -> AsyncIterator[KnowledgeHarness]:
    ks = KnowledgeSettings(**settings.model_dump(), documents_dir=tmp_path / "documents")
    h = KnowledgeHarness(ks)
    async with h.app.router.lifespan_context(h.app):
        yield h
    await h.client.aclose()


@pytest.fixture
async def owner_id(owner: httpx.AsyncClient) -> uuid.UUID:
    return uuid.UUID((await owner.get("/api/v1/me")).json()["user"]["id"])
